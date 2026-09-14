"""Извлечение текста из загруженных документов и разбиение на chunks
(change add-document-rag, design D7/D9).

Чистые функции без ввода-вывода и без сети: на входе имя файла и байты,
на выходе текст или готовые chunks. Все отказы — исключения `DocumentError`
с текстом, пригодным для показа пользователю: обработчик загрузки
пересылает его в чат, а не переводит коды ошибок.
"""
from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import PurePosixPath

SUPPORTED_EXTENSIONS = (".txt", ".md", ".docx", ".pdf")

CHUNK_SIZE = 800
CHUNK_OVERLAP = 120
"""Размер chunk и перекрытие в символах (design D7): ~15% перекрытия, чтобы
факт на стыке окон не терялся. Слишком мелкие chunks теряют контекст абзаца,
слишком крупные размывают релевантность и дороже в эмбеддингах."""


class DocumentError(Exception):
    """Документ нельзя проиндексировать; текст исключения — для пользователя."""


class UnsupportedDocumentType(DocumentError):
    """Расширение файла вне SUPPORTED_EXTENSIONS."""


class DocumentExtractionError(DocumentError):
    """Парсер не смог прочитать файл (повреждён или не тот формат)."""


class EmptyDocumentError(DocumentError):
    """Файл прочитан, но текста в нём нет (пустой или сканированный PDF)."""


class DocumentTooLarge(DocumentError):
    """Документ превышает настроенные лимиты индексации (design D8)."""


@dataclass(frozen=True)
class TextChunk:
    """Фрагмент текста с опциональным span страниц PDF (1-based).

    Для .txt/.md/.docx оба поля страницы — None: номер не выдумывается.
    Для PDF `page_start`/`page_end` указывают страницы, с которых взят
    текст; равны, если фрагмент целиком на одной странице.
    """

    text: str
    page_start: int | None = None
    page_end: int | None = None


def chunk_plain_texts(chunks: list[TextChunk] | list[str]) -> list[str]:
    """Тексты фрагментов — вход эмбеддингов, которым не нужен span страниц."""
    return [chunk.text if isinstance(chunk, TextChunk) else chunk for chunk in chunks]


def document_extension(filename: str) -> str:
    """Расширение файла в нижнем регистре (с точкой) или пустая строка."""
    return PurePosixPath(filename).suffix.lower()


def is_supported_document(filename: str | None) -> bool:
    """Берёт ли RAG такой файл — фильтр обработчика и проверка формата."""
    if not filename:
        return False
    return document_extension(filename) in SUPPORTED_EXTENSIONS


def _supported_list() -> str:
    return ", ".join(SUPPORTED_EXTENSIONS)


def ensure_upload_size(size: int | None, limit: int) -> None:
    """Проверяет сырой размер файла до скачивания и разбора (design D8)."""
    if size is not None and size > limit:
        raise DocumentTooLarge(
            f"файл больше лимита: {size} байт при допустимых {limit}"
        )


def _extract_plain(data: bytes) -> str:
    # Кодировку выбирает пользователь, а не мы: неизвестные байты заменяются,
    # чтобы «почти UTF-8» файл индексировался, а не отвергался целиком.
    return data.decode("utf-8", errors="replace")


def _extract_docx(data: bytes) -> str:
    import docx  # noqa: PLC0415 — тяжёлый парсер грузится только под свой формат

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentExtractionError(
            f"не удалось прочитать .docx: {exc}"
        ) from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    # Политики и гайды часто держат факты в таблицах — иначе они теряются.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cell for cell in cells if cell))
    return "\n".join(part for part in parts if part.strip())


def _extract_pdf_pages(data: bytes) -> tuple[str, tuple[tuple[int, int, int], ...]]:
    """Текст PDF и карта страниц: (номер 1-based, start, end) в этом тексте.

    Пустые страницы пропускаются, но номер берётся из PDF, а не из
    пересчёта непустых. Страницы склеиваются через «\\n», как раньше.
    """
    from pypdf import PdfReader  # noqa: PLC0415

    try:
        reader = PdfReader(io.BytesIO(data))
        raw_pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise DocumentExtractionError(f"не удалось прочитать .pdf: {exc}") from exc
    parts: list[str] = []
    spans: list[tuple[int, int, int]] = []
    offset = 0
    for index, page_text in enumerate(raw_pages):
        if not page_text.strip():
            continue
        if parts:
            offset += 1
        start = offset
        parts.append(page_text)
        offset += len(page_text)
        spans.append((index + 1, start, offset))
    return "\n".join(parts), tuple(spans)


def _extract_pdf(data: bytes) -> str:
    text, _ = _extract_pdf_pages(data)
    return text


def extract_text(filename: str, data: bytes) -> str:
    """Текст документа по его имени и байтам (design D9).

    Формат выбирается по расширению; сканы PDF без текстового слоя дают
    пустой результат и отвергаются как пустой документ (OCR вне скоупа).
    """
    text, _ = extract_document(filename, data)
    return text


def extract_document(
    filename: str, data: bytes
) -> tuple[str, tuple[tuple[int, int, int], ...]]:
    """Текст и карта страниц PDF (пустая для остальных форматов)."""
    extension = document_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentType(
            f"формат {extension or '(без расширения)'} не поддерживается; "
            f"доступны {_supported_list()}"
        )
    if extension == ".docx":
        text, spans = _extract_docx(data), ()
    elif extension == ".pdf":
        text, spans = _extract_pdf_pages(data)
    else:
        text, spans = _extract_plain(data), ()
    if not text.strip():
        raise EmptyDocumentError(
            "в документе не найдено текста (пустой файл или PDF без "
            "текстового слоя — распознавание сканов не поддерживается)"
        )
    return text, spans


def _word_start(text: str, position: int, limit: int) -> int:
    """Первая позиция после пробельного разделителя, начиная с `position`.

    Нужна, чтобы окно начиналось с целого слова: иначе перекрытие сдвигает
    левую границу в середину слова, и chunk открывается огрызком.
    Разделителя до `limit` нет — граница остаётся как есть.
    """
    position = max(0, position)
    candidates = [
        found
        for found in (text.find(" ", position, limit), text.find("\n", position, limit))
        if found != -1
    ]
    return min(candidates) + 1 if candidates else position


def _chunk_windows(
    text: str, size: int, overlap: int
) -> list[tuple[str, int, int]]:
    """Окна chunking: (текст фрагмента, start, end) в уже stripped-тексте."""
    if size <= 0:
        raise ValueError("size must be positive")
    text = text.strip()
    if not text:
        return []
    overlap = max(0, min(overlap, size - 1))
    step = size - overlap
    windows: list[tuple[str, int, int]] = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + size, length)
        if end < length:
            boundary = max(
                text.rfind(" ", start + step, end),
                text.rfind("\n", start + step, end),
            )
            if boundary > start:
                end = boundary
        window = text[start:end]
        chunk = window.strip()
        if chunk:
            rel = window.find(chunk)
            abs_start = start + rel
            windows.append((chunk, abs_start, abs_start + len(chunk)))
        if end >= length:
            break
        start = max(_word_start(text, end - overlap, end), start + 1)
    return windows


def chunk_text(
    text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP
) -> list[str]:
    """Режет текст скользящим окном с перекрытием (design D7).

    Обе границы окна прижимаются к границам слов: правая подтягивается к
    последнему пробелу или переводу строки в пределах последнего сдвига,
    левая — к началу слова внутри перекрытия. Поэтому фактический размер
    окна и перекрытие плавают в пределах заданных, зато chunk не начинается
    и не кончается обрывком слова.

    Правая граница ищется не ближе `size - overlap` от начала окна, а шаг
    дополнительно не даёт start остаться на месте — разбиение конечно при
    любых аргументах, включая overlap больше половины size.
    """
    return [chunk for chunk, _start, _end in _chunk_windows(text, size, overlap)]


def _pages_for_span(
    spans: tuple[tuple[int, int, int], ...], start: int, end: int
) -> tuple[int | None, int | None]:
    """Страницы, пересекающие полуинтервал [start, end) исходного текста."""
    pages = [page for page, page_start, page_end in spans if page_start < end and page_end > start]
    if not pages:
        return None, None
    return pages[0], pages[-1]


def chunks_from_text(
    text: str,
    page_spans: tuple[tuple[int, int, int], ...] = (),
    *,
    max_chars: int,
    max_chunks: int,
    size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[TextChunk]:
    """Лимиты и разбиение уже извлечённого текста (для шагов прогресса)."""
    if len(text) > max_chars:
        raise DocumentTooLarge(
            f"извлечённый текст больше лимита: {len(text)} символов "
            f"при допустимых {max_chars}"
        )
    stripped = text.strip()
    origin = text.find(stripped) if stripped else 0
    windows = _chunk_windows(text, size, overlap)
    if len(windows) > max_chunks:
        raise DocumentTooLarge(
            f"документ даёт {len(windows)} фрагментов при допустимых {max_chunks}"
        )
    if not windows:
        raise EmptyDocumentError("после разбиения не осталось текста")
    chunks: list[TextChunk] = []
    for chunk, start, end in windows:
        page_start = page_end = None
        if page_spans:
            page_start, page_end = _pages_for_span(
                page_spans, origin + start, origin + end
            )
        chunks.append(TextChunk(text=chunk, page_start=page_start, page_end=page_end))
    return chunks


def split_document(
    filename: str,
    data: bytes,
    *,
    max_chars: int,
    max_chunks: int,
    size: int = CHUNK_SIZE,
    overlap: int = CHUNK_OVERLAP,
) -> list[TextChunk]:
    """Полный путь от байтов до chunks с проверкой лимитов (design D8).

    Лимиты объёма текста и числа chunks проверяются здесь, после извлечения
    и разбиения: до разбора их значения неизвестны. Исключение означает, что
    документ не индексируется вовсе — частичного индекса не остаётся.
    Для PDF у фрагмента заполнены page_start/page_end; иначе оба None.
    """
    text, page_spans = extract_document(filename, data)
    return chunks_from_text(
        text,
        page_spans,
        max_chars=max_chars,
        max_chunks=max_chunks,
        size=size,
        overlap=overlap,
    )
